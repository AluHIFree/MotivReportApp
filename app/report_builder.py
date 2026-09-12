# -*- coding: utf-8 -*-
"""Построение сводного Excel-отчёта мотивирующего мониторинга."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .theme import DEFAULT_OUTPUT_NAME, REQUIRED_COLUMNS

NAVY = "1A4C8C"
NAVY2 = "2563A8"
ALT = "F8FAFC"
WHITE = "FFFFFF"
GREEN = "059669"
GREEN_BG = "D1FAE5"
AMBER = "D97706"
AMBER_BG = "FEF3C7"
RED = "DC2626"
RED_BG = "FEE2E2"
DARK = "0F172A"
GRAY = "64748B"
BORDER_C = "CBD5E1"
TEAL = "0F766E"


class ReportValidationError(ValueError):
    """Ошибка проверки входных выгрузок."""


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _score_fill(v: float | None) -> PatternFill:
    if v is None:
        return _fill("E2E8F0")
    if v >= 95:
        return _fill(GREEN_BG)
    if v >= 75:
        return _fill(AMBER_BG)
    return _fill(RED_BG)


def _score_font(v: float | None) -> Font:
    if v is None:
        return Font(name="Calibri", size=11, color=GRAY)
    if v >= 95:
        return Font(name="Calibri", size=11, bold=True, color=GREEN)
    if v >= 75:
        return Font(name="Calibri", size=11, bold=True, color=AMBER)
    return Font(name="Calibri", size=11, bold=True, color=RED)


def _err_fill(n: int) -> PatternFill:
    if n == 0:
        return _fill(GREEN_BG)
    if n <= 10:
        return _fill(AMBER_BG)
    return _fill(RED_BG)


def _err_font(n: int) -> Font:
    if n == 0:
        return Font(name="Calibri", size=11, bold=True, color=GREEN)
    if n <= 10:
        return Font(name="Calibri", size=11, bold=True, color=AMBER)
    return Font(name="Calibri", size=11, bold=True, color=RED)


def _thin_border() -> Border:
    side = Side(style="thin", color=BORDER_C)
    return Border(left=side, right=side, top=side, bottom=side)


def validate_export(path: str | Path, kind: str) -> pd.DataFrame:
    """Читает Excel и проверяет обязательные столбцы."""
    path = Path(path)
    if not path.exists():
        raise ReportValidationError(f"Файл не найден: {path.name}")
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        raise ReportValidationError(f"Ожидается Excel-файл (.xlsx): {path.name}")

    try:
        df = pd.read_excel(path)
    except Exception as exc:  # noqa: BLE001
        raise ReportValidationError(f"Не удалось прочитать «{path.name}»: {exc}") from exc

    if df.empty:
        raise ReportValidationError(f"Файл «{path.name}» пустой")

    required = REQUIRED_COLUMNS[kind]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ReportValidationError(
            f"В файле «{path.name}» нет столбцов: {', '.join(missing)}"
        )
    return df


def _grade_percent(v: float) -> str:
    if v >= 95:
        return "Высокий"
    if v >= 75:
        return "Средний"
    return "Низкий"


def _grade_errors(n: int) -> str:
    if n <= 20:
        return "Низкий риск"
    if n <= 50:
        return "Средний риск"
    return "Высокий риск"


def build_summary_report(
    tim_path: str | Path,
    tm_path: str | Path,
    err_path: str | Path,
    out_path: str | Path | None = None,
) -> dict[str, Any]:
    """Строит сводный xlsx и возвращает метрики для UI."""
    tim = validate_export(tim_path, "timeliness")
    tm = validate_export(tm_path, "typical_menu")
    err = validate_export(err_path, "sanpin")

    tim = tim.copy()
    tm = tm.copy()
    err = err.copy()

    tim["Процент Своевременности"] = pd.to_numeric(
        tim["Процент Своевременности"], errors="coerce"
    )
    tm["Процент Соблюдения"] = pd.to_numeric(tm["Процент Соблюдения"], errors="coerce")

    if tim["Процент Своевременности"].isna().all():
        raise ReportValidationError("В выгрузке своевременности нет числовых процентов")
    if tm["Процент Соблюдения"].isna().all():
        raise ReportValidationError("В выгрузке соответствия ТМ нет числовых процентов")

    districts = sorted(set(tim["Район"]).union(tm["Район"]).union(err["Район"]))

    tim_d = tim.groupby("Район").agg(
        tim_mean=("Процент Своевременности", "mean"),
        tim_n=("Процент Своевременности", "count"),
        tim_100=("Процент Своевременности", lambda s: int((s == 100).sum())),
        tim_75=("Процент Своевременности", lambda s: int((s >= 75).sum())),
    )
    tm_d = tm.groupby("Район").agg(
        tm_mean=("Процент Соблюдения", "mean"),
        tm_n=("Процент Соблюдения", "count"),
        tm_100=("Процент Соблюдения", lambda s: int((s == 100).sum())),
        tm_75=("Процент Соблюдения", lambda s: int((s >= 75).sum())),
        tm_fac=("Пищеблок Наименование", "nunique"),
    )
    err_d = err.groupby("Район").agg(
        err_n=("Ошибка", "count"),
        err_fac=("Пищеблок", "nunique"),
    )

    rows: list[dict[str, Any]] = []
    for d in districts:
        t = tim_d.loc[d] if d in tim_d.index else None
        m = tm_d.loc[d] if d in tm_d.index else None
        e = err_d.loc[d] if d in err_d.index else None
        tim_mean = float(t.tim_mean) if t is not None else None
        tm_mean = float(m.tm_mean) if m is not None else None
        err_n = int(e.err_n) if e is not None else 0
        if tim_mean is not None and tm_mean is not None:
            composite = (tim_mean + tm_mean) / 2
        elif tim_mean is not None:
            composite = tim_mean
        elif tm_mean is not None:
            composite = tm_mean
        else:
            composite = None
        rows.append(
            {
                "district": d,
                "tim_mean": tim_mean,
                "tim_n": int(t.tim_n) if t is not None else 0,
                "tim_100": int(t.tim_100) if t is not None else 0,
                "tim_75": int(t.tim_75) if t is not None else 0,
                "tm_mean": tm_mean,
                "tm_n": int(m.tm_n) if m is not None else 0,
                "tm_100": int(m.tm_100) if m is not None else 0,
                "tm_75": int(m.tm_75) if m is not None else 0,
                "tm_fac": int(m.tm_fac) if m is not None else 0,
                "err_n": err_n,
                "err_fac": int(e.err_fac) if e is not None else 0,
                "composite": composite,
            }
        )

    df = (
        pd.DataFrame(rows)
        .sort_values(["composite", "err_n"], ascending=[False, True], na_position="last")
        .reset_index(drop=True)
    )

    reg_tim_mean = float(tim["Процент Своевременности"].mean())
    reg_tm_mean = float(tm["Процент Соблюдения"].mean())
    reg_tim_100_pct = float((tim["Процент Своевременности"] == 100).mean() * 100)
    reg_tm_100_pct = float((tm["Процент Соблюдения"] == 100).mean() * 100)
    reg_tim_75_pct = float((tim["Процент Своевременности"] >= 75).mean() * 100)
    reg_tm_75_pct = float((tm["Процент Соблюдения"] >= 75).mean() * 100)
    reg_err = len(err)
    reg_composite = (reg_tim_mean + reg_tm_mean) / 2
    tim_100_n = int((tim["Процент Своевременности"] == 100).sum())
    tim_75_n = int((tim["Процент Своевременности"] >= 75).sum())
    tm_100_n = int((tm["Процент Соблюдения"] == 100).sum())
    tm_75_n = int((tm["Процент Соблюдения"] >= 75).sum())
    tm_facilities = int(tm["Пищеблок Наименование"].nunique())
    err_facilities = int(err["Пищеблок"].nunique())
    err_districts = int(err["Район"].nunique())

    if out_path is None:
        out_path = Path(tm_path).resolve().parent / DEFAULT_OUTPUT_NAME
    else:
        out_path = Path(out_path)
        if out_path.is_dir():
            out_path = out_path / DEFAULT_OUTPUT_NAME
        if out_path.suffix.lower() != ".xlsx":
            out_path = out_path.with_suffix(".xlsx")

    out_path.parent.mkdir(parents=True, exist_ok=True)

    thin = _thin_border()
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    wb = Workbook()

    # ----- Sheet 1: region -----
    ws = wb.active
    ws.title = "Сводка по региону"

    ws.merge_cells("B2:G2")
    ws["B2"] = "Мотивирующий мониторинг — сводный отчёт"
    ws["B2"].font = Font(name="Calibri", size=20, bold=True, color=NAVY)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("B3:G3")
    region_name = str(tim["Регион"].dropna().iloc[0]) if tim["Регион"].notna().any() else "Регион"
    ws["B3"] = (
        f"{region_name}  •  период данных по выгрузкам  •  сформировано: "
        + datetime.now().strftime("%d.%m.%Y %H:%M")
    )
    ws["B3"].font = Font(name="Calibri", size=11, color=GRAY)

    ws.merge_cells("B5:G5")
    ws["B5"] = "ПОКАЗАТЕЛИ ПО РЕГИОНУ В ЦЕЛОМ"
    ws["B5"].font = Font(name="Calibri", size=12, bold=True, color=WHITE)
    ws["B5"].fill = _fill(NAVY)
    ws["B5"].alignment = center
    for col in range(2, 8):
        ws.cell(5, col).fill = _fill(NAVY)

    headers = ["Показатель", "Значение", "Охват / объём", "100% выполнения", "≥75%", "Оценка"]
    for i, h in enumerate(headers, start=2):
        cell = ws.cell(6, i, h)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = _fill(NAVY2)
        cell.alignment = center
        cell.border = thin

    kpi_rows = [
        [
            "Своевременность размещения меню",
            f"{reg_tim_mean:.2f}%",
            f"{len(tim)} пищеблоков",
            f"{reg_tim_100_pct:.2f}% ({tim_100_n} из {len(tim)})",
            f"{reg_tim_75_pct:.2f}% ({tim_75_n})",
            _grade_percent(reg_tim_mean),
            reg_tim_mean,
        ],
        [
            "Соответствие типовому меню",
            f"{reg_tm_mean:.2f}%",
            f"{len(tm)} записей / {tm_facilities} пищеблоков",
            f"{reg_tm_100_pct:.2f}% ({tm_100_n} из {len(tm)})",
            f"{reg_tm_75_pct:.2f}% ({tm_75_n})",
            _grade_percent(reg_tm_mean),
            reg_tm_mean,
        ],
        [
            "Ошибки по СанПиН",
            f"{reg_err} шт.",
            f"{err_facilities} пищеблоков с ошибками / {err_districts} районов",
            "—",
            "—",
            _grade_errors(reg_err),
            None,
        ],
    ]

    for r_idx, row in enumerate(kpi_rows, start=7):
        for c_idx, val in enumerate(row[:6], start=2):
            cell = ws.cell(r_idx, c_idx, val)
            cell.border = thin
            cell.alignment = center if c_idx > 2 else left
            cell.font = Font(name="Calibri", size=11, bold=(c_idx == 2))
            cell.fill = _fill(ALT if r_idx % 2 == 0 else WHITE)
        score = row[6]
        if score is not None:
            ws.cell(r_idx, 3).fill = _score_fill(score)
            ws.cell(r_idx, 3).font = _score_font(score)
            ws.cell(r_idx, 7).fill = _score_fill(score)
            ws.cell(r_idx, 7).font = _score_font(score)
        else:
            ws.cell(r_idx, 3).fill = _err_fill(reg_err)
            ws.cell(r_idx, 3).font = _err_font(reg_err)
            ws.cell(r_idx, 7).fill = _err_fill(reg_err)
            ws.cell(r_idx, 7).font = _err_font(reg_err)

    ws.merge_cells("B11:G11")
    ws["B11"] = "ИТОГОВОЕ ПОКРЫТИЕ РЕГИОНА"
    ws["B11"].font = Font(name="Calibri", size=12, bold=True, color=WHITE)
    ws["B11"].fill = _fill(TEAL)
    ws["B11"].alignment = center
    for col in range(2, 8):
        ws.cell(11, col).fill = _fill(TEAL)

    top3 = ", ".join(df.head(3)["district"].tolist())
    bottom3 = ", ".join(df.tail(3)["district"].tolist())
    err_top = ", ".join(
        df.sort_values("err_n", ascending=False)
        .head(3)
        .apply(lambda r: f"{r.district} ({r.err_n})", axis=1)
        .tolist()
    )
    cover_text = (
        f"Среднее по двум процентным показателям (своевременность + соответствие ТМ): "
        f"{reg_composite:.2f}%.\n"
        f"Своевременность: {reg_tim_mean:.2f}% по {len(tim)} пищеблокам. "
        f"Соответствие типовому меню: {reg_tm_mean:.2f}% по {len(tm)} дневным записям "
        f"({tm_facilities} пищеблоков). "
        f"Выявлено ошибок СанПиН: {reg_err} "
        f"(в {err_districts} районах, {err_facilities} пищеблоков).\n"
        f"Районов в выборке: {len(districts)}. Лучшие по комплексному %: {top3}. "
        f"Требуют внимания: {bottom3}. "
        f"Наибольшее число ошибок СанПиН: {err_top}."
    )

    ws.merge_cells("B12:G14")
    ws["B12"] = cover_text
    ws["B12"].font = Font(name="Calibri", size=11, color=DARK)
    ws["B12"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["B12"].fill = _fill("ECFDF5")
    for col in range(2, 8):
        for row in range(12, 15):
            ws.cell(row, col).fill = _fill("ECFDF5")
            ws.cell(row, col).border = thin

    ws.merge_cells("B16:G16")
    ws["B16"] = (
        "Шкала оценки процентных показателей: ≥95% — высокий (зелёный)  |  "
        "75–94% — средний (янтарный)  |  <75% — низкий (красный).  "
        "СанПиН: 0 — без ошибок; 1–10 — умеренно; >10 — много ошибок."
    )
    ws["B16"].font = Font(name="Calibri", size=9, italic=True, color=GRAY)
    ws["B16"].alignment = left

    ws.merge_cells("B18:G18")
    ws["B18"] = (
        f"Источники: {Path(tim_path).name}; {Path(tm_path).name}; {Path(err_path).name}"
    )
    ws["B18"].font = Font(name="Calibri", size=9, color=GRAY)

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 42
    ws.column_dimensions["E"].width = 36
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 14
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[5].height = 22
    ws.row_dimensions[6].height = 30
    for r in range(7, 10):
        ws.row_dimensions[r].height = 28
    ws.row_dimensions[12].height = 28
    ws.row_dimensions[13].height = 28
    ws.row_dimensions[14].height = 28
    ws.sheet_view.showGridLines = False

    # ----- Sheet 2: districts -----
    ws2 = wb.create_sheet("По районам")
    ws2.merge_cells("B2:N2")
    ws2["B2"] = "Статистика мотивирующего мониторинга в разрезе районов"
    ws2["B2"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)

    ws2.merge_cells("B3:N3")
    ws2["B3"] = (
        "Своевременность размещения меню • Соответствие типовому меню • Ошибки по СанПиН"
    )
    ws2["B3"].font = Font(name="Calibri", size=11, color=GRAY)

    ws2.merge_cells("B5:B6")
    ws2["B5"] = "№"
    ws2.merge_cells("C5:C6")
    ws2["C5"] = "Район"
    ws2.merge_cells("D5:G5")
    ws2["D5"] = "Своевременность размещения меню"
    ws2.merge_cells("H5:K5")
    ws2["H5"] = "Соответствие типовому меню"
    ws2.merge_cells("L5:M5")
    ws2["L5"] = "СанПиН"
    ws2.merge_cells("N5:N6")
    ws2["N5"] = "Комплексный\nпоказатель, %"

    sub = {
        "D6": "Средний %",
        "E6": "Пищеблоков",
        "F6": "100%",
        "G6": "≥75%",
        "H6": "Средний %",
        "I6": "Записей",
        "J6": "100%",
        "K6": "≥75%",
        "L6": "Ошибок",
        "M6": "Пищеблоков\nс ошибками",
    }
    for addr, text in sub.items():
        ws2[addr] = text

    for addr in [
        "B5", "C5", "D5", "E5", "F5", "G5", "H5", "I5", "J5", "K5", "L5", "M5", "N5",
        "B6", "C6", "D6", "E6", "F6", "G6", "H6", "I6", "J6", "K6", "L6", "M6", "N6",
    ]:
        cell = ws2[addr]
        cell.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
        cell.alignment = center
        cell.border = thin

    for col in range(2, 4):
        for row in (5, 6):
            ws2.cell(row, col).fill = _fill(NAVY)
    for col in range(4, 8):
        for row in (5, 6):
            ws2.cell(row, col).fill = _fill("1D4ED8")
    for col in range(8, 12):
        for row in (5, 6):
            ws2.cell(row, col).fill = _fill("0F766E")
    for col in range(12, 14):
        for row in (5, 6):
            ws2.cell(row, col).fill = _fill("B45309")
    for row in (5, 6):
        ws2.cell(row, 14).fill = _fill("6D28D9")

    for i, r in df.iterrows():
        rr = 7 + i
        values = [
            i + 1,
            r["district"],
            round(r["tim_mean"], 2) if pd.notna(r["tim_mean"]) else None,
            r["tim_n"],
            r["tim_100"],
            r["tim_75"],
            round(r["tm_mean"], 2) if pd.notna(r["tm_mean"]) else None,
            r["tm_n"],
            r["tm_100"],
            r["tm_75"],
            r["err_n"],
            r["err_fac"],
            round(r["composite"], 2) if pd.notna(r["composite"]) else None,
        ]
        for c_idx, val in enumerate(values, start=2):
            cell = ws2.cell(rr, c_idx, val)
            cell.border = thin
            cell.alignment = center if c_idx != 3 else left
            cell.font = Font(name="Calibri", size=10)
            if i % 2 == 1:
                cell.fill = _fill(ALT)

        for col, key in [(4, "tim_mean"), (8, "tm_mean"), (14, "composite")]:
            v = r[key]
            cell = ws2.cell(rr, col)
            if pd.notna(v):
                cell.number_format = "0.00"
                cell.fill = _score_fill(float(v))
                cell.font = _score_font(float(v))
            else:
                cell.value = "—"

        ws2.cell(rr, 12).fill = _err_fill(int(r["err_n"]))
        ws2.cell(rr, 12).font = _err_font(int(r["err_n"]))
        ws2.row_dimensions[rr].height = 20

    last_data = 6 + len(df)
    tot_row = last_data + 1
    ws2.cell(tot_row, 2, "")
    ws2.cell(tot_row, 3, "РЕГИОН ИТОГО")
    ws2.cell(tot_row, 4, round(reg_tim_mean, 2))
    ws2.cell(tot_row, 5, len(tim))
    ws2.cell(tot_row, 6, tim_100_n)
    ws2.cell(tot_row, 7, tim_75_n)
    ws2.cell(tot_row, 8, round(reg_tm_mean, 2))
    ws2.cell(tot_row, 9, len(tm))
    ws2.cell(tot_row, 10, tm_100_n)
    ws2.cell(tot_row, 11, tm_75_n)
    ws2.cell(tot_row, 12, reg_err)
    ws2.cell(tot_row, 13, err_facilities)
    ws2.cell(tot_row, 14, round(reg_composite, 2))

    for col in range(2, 15):
        cell = ws2.cell(tot_row, col)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = _fill(NAVY)
        cell.alignment = center
        cell.border = thin
    ws2.cell(tot_row, 3).alignment = left
    ws2.cell(tot_row, 4).number_format = "0.00"
    ws2.cell(tot_row, 8).number_format = "0.00"
    ws2.cell(tot_row, 14).number_format = "0.00"
    ws2.row_dimensions[tot_row].height = 24

    note_row = tot_row + 2
    ws2.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=14)
    ws2.cell(
        note_row,
        2,
        "Пояснения: «Средний %» — среднее арифметическое по пищеблокам/записям района. "
        "«100%» / «≥75%» — число единиц с соответствующим уровнем. "
        "«Комплексный показатель» — среднее своевременности и соответствия ТМ. "
        "Районы без ошибок СанПиН отображаются с нулём.",
    )
    ws2.cell(note_row, 2).font = Font(name="Calibri", size=9, italic=True, color=GRAY)

    ws2.column_dimensions["A"].width = 2
    ws2.column_dimensions["B"].width = 5
    ws2.column_dimensions["C"].width = 22
    for col in range(4, 15):
        ws2.column_dimensions[get_column_letter(col)].width = 12
    ws2.column_dimensions["N"].width = 14
    ws2.row_dimensions[2].height = 26
    ws2.row_dimensions[5].height = 22
    ws2.row_dimensions[6].height = 32
    ws2.freeze_panes = "D7"
    if len(df) > 0:
        ws2.auto_filter.ref = f"B6:N{tot_row - 1}"
    ws2.sheet_view.showGridLines = False

    # ----- Sheet 3: SanPiN -----
    ws3 = wb.create_sheet("СанПиН по районам")
    ws3.merge_cells("B2:F2")
    ws3["B2"] = "Детализация ошибок по СанПиН"
    ws3["B2"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)

    ws3.merge_cells("B3:F3")
    ws3["B3"] = (
        f"Всего ошибок: {reg_err}  •  пищеблоков с ошибками: {err_facilities}  "
        f"•  районов: {err_districts}"
    )
    ws3["B3"].font = Font(name="Calibri", size=11, color=GRAY)

    err_summary = (
        err.groupby("Район")
        .agg(
            ошибок=("Ошибка", "count"),
            пищеблоков=("Пищеблок", "nunique"),
            видов_ошибок=("Ошибка", "nunique"),
        )
        .reset_index()
        .sort_values("ошибок", ascending=False)
        .reset_index(drop=True)
    )
    err_types = err["Ошибка"].value_counts().reset_index()
    err_types.columns = ["Тип ошибки", "Количество"]

    headers3 = ["№", "Район", "Кол-во ошибок", "Пищеблоков с ошибками", "Уникальных типов ошибок"]
    for i, h in enumerate(headers3, start=2):
        cell = ws3.cell(5, i, h)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = _fill("B45309")
        cell.alignment = center
        cell.border = thin

    for i, row in err_summary.iterrows():
        rr = 6 + i
        vals = [
            i + 1,
            row["Район"],
            int(row["ошибок"]),
            int(row["пищеблоков"]),
            int(row["видов_ошибок"]),
        ]
        for c, v in enumerate(vals, start=2):
            cell = ws3.cell(rr, c, v)
            cell.border = thin
            cell.alignment = center if c != 3 else left
            cell.font = Font(name="Calibri", size=10)
            if c == 4:
                cell.fill = _err_fill(int(v))
                cell.font = _err_font(int(v))
            elif i % 2 == 1:
                cell.fill = _fill(ALT)

    zero = [d for d in districts if d not in set(err["Район"])]
    zr = 6 + len(err_summary) + 1
    ws3.merge_cells(start_row=zr, start_column=2, end_row=zr, end_column=6)
    ws3.cell(zr, 2, "Районы без ошибок СанПиН: " + (", ".join(zero) if zero else "—"))
    ws3.cell(zr, 2).font = Font(name="Calibri", size=10, color=GREEN)
    ws3.cell(zr, 2).fill = _fill(GREEN_BG)

    tr = zr + 2
    ws3.merge_cells(start_row=tr, start_column=2, end_row=tr, end_column=4)
    ws3.cell(tr, 2, "Топ типов ошибок СанПиН")
    ws3.cell(tr, 2).font = Font(name="Calibri", size=12, bold=True, color=NAVY)

    for i, h in enumerate(["№", "Тип ошибки", "Количество"], start=2):
        cell = ws3.cell(tr + 1, i, h)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = _fill(NAVY)
        cell.alignment = center
        cell.border = thin

    for i, row in err_types.iterrows():
        rr = tr + 2 + i
        for c, v in enumerate([i + 1, row["Тип ошибки"], int(row["Количество"])], start=2):
            cell = ws3.cell(rr, c, v)
            cell.border = thin
            cell.alignment = left if c == 3 else center
            cell.font = Font(name="Calibri", size=10)
            if i % 2 == 1:
                cell.fill = _fill(ALT)

    ws3.column_dimensions["A"].width = 2
    ws3.column_dimensions["B"].width = 5
    ws3.column_dimensions["C"].width = 55
    ws3.column_dimensions["D"].width = 22
    ws3.column_dimensions["E"].width = 26
    ws3.column_dimensions["F"].width = 14
    ws3.sheet_view.showGridLines = False

    # ----- Sheet 4: ranking -----
    ws4 = wb.create_sheet("Рейтинг районов")
    ws4.merge_cells("B2:G2")
    ws4["B2"] = "Рейтинг районов по комплексному покрытию показателей"
    ws4["B2"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)

    headers4 = [
        "Место",
        "Район",
        "Своевременность, %",
        "Соответствие ТМ, %",
        "Ошибки СанПиН",
        "Комплексный %",
    ]
    for i, h in enumerate(headers4, start=2):
        cell = ws4.cell(4, i, h)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = _fill(NAVY)
        cell.alignment = center
        cell.border = thin

    medal = ["FBBF24", "CBD5E1", "D97706"]
    for i, r in df.iterrows():
        rr = 5 + i
        vals = [
            i + 1,
            r["district"],
            round(r["tim_mean"], 2) if pd.notna(r["tim_mean"]) else None,
            round(r["tm_mean"], 2) if pd.notna(r["tm_mean"]) else None,
            int(r["err_n"]),
            round(r["composite"], 2) if pd.notna(r["composite"]) else None,
        ]
        for c, v in enumerate(vals, start=2):
            cell = ws4.cell(rr, c, v if v is not None else "—")
            cell.border = thin
            cell.alignment = center if c != 3 else left
            cell.font = Font(name="Calibri", size=11)
            if i % 2 == 1 and c not in (4, 5, 6, 7):
                cell.fill = _fill(ALT)
        for col, key in [(4, "tim_mean"), (5, "tm_mean"), (7, "composite")]:
            v = r[key]
            if pd.notna(v):
                ws4.cell(rr, col).number_format = "0.00"
                ws4.cell(rr, col).fill = _score_fill(float(v))
                ws4.cell(rr, col).font = _score_font(float(v))
        ws4.cell(rr, 6).fill = _err_fill(int(r["err_n"]))
        ws4.cell(rr, 6).font = _err_font(int(r["err_n"]))
        if i < 3:
            ws4.cell(rr, 2).fill = _fill(medal[i])
            ws4.cell(rr, 2).font = Font(name="Calibri", size=11, bold=True, color=DARK)
        ws4.row_dimensions[rr].height = 22

    ws4.column_dimensions["A"].width = 2
    ws4.column_dimensions["B"].width = 8
    ws4.column_dimensions["C"].width = 22
    ws4.column_dimensions["D"].width = 20
    ws4.column_dimensions["E"].width = 20
    ws4.column_dimensions["F"].width = 16
    ws4.column_dimensions["G"].width = 16
    ws4.freeze_panes = "C5"
    ws4.sheet_view.showGridLines = False

    try:
        wb.save(out_path)
    except PermissionError as exc:
        raise ReportValidationError(
            f"Не удалось сохранить отчёт — файл открыт в другой программе:\n{out_path}"
        ) from exc
    except OSError as exc:
        raise ReportValidationError(f"Ошибка записи отчёта: {exc}") from exc

    return {
        "output_path": str(out_path.resolve()),
        "region": region_name,
        "districts": len(districts),
        "timeliness_mean": round(reg_tim_mean, 2),
        "typical_menu_mean": round(reg_tm_mean, 2),
        "composite": round(reg_composite, 2),
        "sanpin_errors": reg_err,
        "timeliness_count": len(tim),
        "typical_menu_records": len(tm),
        "typical_menu_facilities": tm_facilities,
        "top_districts": df.head(3)["district"].tolist(),
        "attention_districts": df.tail(3)["district"].tolist(),
    }
